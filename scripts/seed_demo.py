"""One-time demo data seed for VTI ERP (Firestore project skillbridge-3b6c3).

Schema is taken from the app's screens:
  batches     -> lib/add_batch_screen.dart
  students    -> lib/add_student_screen.dart (+ counters/traineeId_<year>)
  attendance  -> lib/mark_attendance_screen.dart (doc id = <batchId>_<yyyy-mm-dd>)
  companies   -> lib/add_company_screen.dart
  placements  -> lib/add_placement_screen.dart
Every seeded record gets demoSeed: true.

Credentials: put the service account key path in the GOOGLE_APPLICATION_CREDENTIALS env var.

Usage:
  python scripts/seed_demo.py --excel path/to/demo_trainees_Vistas.xlsx --dry-run
  python scripts/seed_demo.py --excel path/to/demo_trainees_Vistas.xlsx
  python scripts/seed_demo.py --clear [--dry-run]
"""

import argparse
import os
import random
import re
import sys
from datetime import date, datetime, timedelta, timezone

PROJECT_ID = 'skillbridge-3b6c3'
YEAR = 2026
COUNTER_DOC = f'traineeId_{YEAR}'
IST = timezone(timedelta(hours=5, minutes=30))
TODAY = datetime.now(IST).date()
BATCH_WRITE_LIMIT = 400
EXPECTED_TRAINEES = 298
TARGET_PLACEMENTS = 120

# batchName -> (course, startDate)
BATCHES = {
    'ELC-01': ('Electrical', date(2026, 3, 2)),
    'ELC-02': ('Electrical', date(2026, 7, 6)),
    'BPO-01': ('BPO', date(2026, 3, 9)),
    'BPO-02': ('BPO', date(2026, 7, 13)),
    'RTL-01': ('Retail', date(2026, 3, 16)),
    'RTL-02': ('Retail', date(2026, 8, 3)),
    'BTY-01': ('Beauty', date(2026, 6, 1)),
    'CAD-01': ('CAD', date(2026, 3, 23)),
    'CAD-02': ('CAD', date(2026, 8, 10)),
    'HSP-01': ('Hospitality', date(2026, 7, 20)),
}

# Durations (months) from lib/course_data.dart
COURSE_MONTHS = {'Electrical': 4, 'BPO': 4, 'Retail': 3, 'Beauty': 4, 'CAD': 4, 'Hospitality': 3}
BATCH_TIMING = '10:00 AM - 2:00 PM'

STUDENT_STATUSES = ['Registered', 'Joined', 'Ongoing', 'Completed', 'Dropped']
NO_ATTENDANCE_STATUSES = {'Registered', 'Dropped'}

# Fictional companies: (name, contactPerson, contactNumber, location, courses)
COMPANIES = [
    ('Shakti Power Solutions', 'Ramesh Verma', '9826011201', 'Indore', ['Electrical']),
    ('Narmada Electricals Pvt Ltd', 'Sunil Patidar', '9826011202', 'Bhopal', ['Electrical', 'CAD']),
    ('Vindhya Call Services', 'Pooja Sharma', '9826011203', 'Indore', ['BPO']),
    ('Malwa Connect BPO', 'Anil Jain', '9826011204', 'Ujjain', ['BPO', 'Retail']),
    ('Apna Bazaar Retail', 'Kavita Rathore', '9826011205', 'Dewas', ['Retail']),
    ('Shree Mart Supermarket', 'Deepak Agrawal', '9826011206', 'Bhopal', ['Retail', 'Hospitality']),
    ('Glow & Grace Salon', 'Neha Tiwari', '9826011207', 'Indore', ['Beauty']),
    ('Roop Shringar Beauty Studio', 'Sapna Chouhan', '9826011208', 'Ujjain', ['Beauty']),
    ('DesignLine Drafting Studio', 'Vikas Dubey', '9826011209', 'Pune', ['CAD']),
    ('Sagar Infra Consultants', 'Manoj Yadav', '9826011210', 'Indore', ['CAD', 'Electrical']),
    ('Hotel Rajwada Residency', 'Arvind Singh', '9826011211', 'Indore', ['Hospitality']),
]

JOB_POSTS = {
    'Electrical': ['Assistant Electrician', 'Wireman', 'Maintenance Helper'],
    'BPO': ['Customer Care Executive', 'Telecaller', 'Back Office Associate'],
    'Retail': ['Retail Sales Associate', 'Cashier', 'Store Assistant'],
    'Beauty': ['Beauty Assistant', 'Salon Assistant', 'Beautician'],
    'CAD': ['CAD Draftsman', 'Junior Draughtsman', 'CAD Operator'],
    'Hospitality': ['Front Office Assistant', 'Housekeeping Attendant', 'F&B Steward'],
}

SEED_COLLECTIONS = ['attendance', 'placements', 'companies', 'students', 'batches']


# ---------- helpers ----------

def add_months(d, months):
    m = d.month - 1 + months
    y = d.year + m // 12
    m = m % 12 + 1
    days_in_month = [31, 29 if y % 4 == 0 and (y % 100 != 0 or y % 400 == 0) else 28,
                     31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]
    return date(y, m, min(d.day, days_in_month))


def ist_midnight(d):
    """The app saves the local (IST) midnight DateTime; same here."""
    return datetime(d.year, d.month, d.day, tzinfo=IST)


def norm_header(v):
    return re.sub(r'[^a-z]', '', str(v or '').lower())


def norm_batch(v):
    s = re.sub(r'[\s_]+', '-', str(v or '').strip().upper())
    m = re.match(r'^([A-Z]+)-?(\d+)$', s)
    return f'{m.group(1)}-{int(m.group(2)):02d}' if m else s


def cell_str(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def cell_num(v):
    if v is None or v == '':
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r'[^\d.]', '', str(v))
    return float(s) if s else 0.0


def cell_date(v):
    if v is None or v == '':
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%d.%m.%Y'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def norm_status(v):
    s = cell_str(v).lower()
    for opt in STUDENT_STATUSES:
        if opt.lower() == s:
            return opt
    return None


def norm_phone(v):
    return re.sub(r'\D', '', cell_str(v))


# ---------- Excel ----------

HEADER_KEYS = {
    'srno': 'srno', 'gender': 'gender', 'name': 'name', 'batch': 'batch', 'dob': 'dob',
    'village': 'village', 'tehsil': 'tehsil', 'contact': 'contact', 'emergency': 'emergency',
    'feesreceivedamount': 'feePaid', 'receiptno': 'receiptNo', 'feesreceiveddate': 'feeDate',
    'totalfee': 'totalFee', 'status': 'status',
}
REQUIRED_COLS = ['gender', 'name', 'batch', 'dob', 'village', 'tehsil', 'contact',
                 'emergency', 'feePaid', 'totalFee', 'status']


def read_trainees(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    trainees, warnings = [], []
    for ws in wb.worksheets:
        rows = ws.iter_rows(min_row=2, values_only=True)
        header = next(rows, None)
        if header is None:
            warnings.append(f'[{ws.title}] khali sheet')
            continue
        col = {}
        for i, h in enumerate(header):
            key = HEADER_KEYS.get(norm_header(h))
            if key and key not in col:
                col[key] = i
        missing = [c for c in REQUIRED_COLS if c not in col]
        if missing:
            raise SystemExit(f'[{ws.title}] columns missing from header: {missing}')

        for r, row in enumerate(rows, start=3):
            get = lambda k: row[col[k]] if col[k] < len(row) else None
            name = cell_str(get('name'))
            if not name:
                continue
            where = f'[{ws.title} row {r}] {name}'
            batch = norm_batch(get('batch')) or norm_batch(ws.title)
            if batch not in BATCHES:
                warnings.append(f'{where}: batch "{batch}" not in list, skipped')
                continue
            status = norm_status(get('status'))
            if status is None:
                warnings.append(f'{where}: status "{cell_str(get("status"))}" not in app options, '
                                f'treated as "Registered"')
                status = 'Registered'
            g = cell_str(get('gender')).upper()[:1]
            gender = {'M': 'Male', 'F': 'Female'}.get(g, 'Other')
            dob = cell_date(get('dob'))
            if get('dob') not in (None, '') and dob is None:
                warnings.append(f'{where}: DOB "{get("dob")}" could not be parsed, left empty')
            total_fee = cell_num(get('totalFee'))
            fee_paid = cell_num(get('feePaid'))
            if fee_paid > total_fee:
                warnings.append(f'{where}: fee paid {fee_paid} > total {total_fee}, set equal to total')
                fee_paid = total_fee
            trainees.append({
                'name': name,
                'phone': norm_phone(get('contact')),
                'gender': gender,
                'dob': dob,
                'village': cell_str(get('village')),
                'tehsil': cell_str(get('tehsil')),
                'guardianPhone': norm_phone(get('emergency')),
                'batch': batch,
                'status': status,
                'totalFee': total_fee,
                'feePaid': fee_paid,
            })
    wb.close()
    return trainees, warnings


# ---------- Firestore ----------

def init_db():
    cred_path = os.environ.get('GOOGLE_APPLICATION_CREDENTIALS')
    if not cred_path:
        raise SystemExit('GOOGLE_APPLICATION_CREDENTIALS env var is not set.')
    if not os.path.isfile(cred_path):
        raise SystemExit('GOOGLE_APPLICATION_CREDENTIALS file not found.')

    import firebase_admin
    from firebase_admin import credentials, firestore

    firebase_admin.initialize_app(credentials.ApplicationDefault(), {'projectId': PROJECT_ID})
    return firestore.client()


class Writer:
    """Batched writes (400 per commit); dry-run only counts."""

    def __init__(self, db, dry_run):
        self.db, self.dry_run = db, dry_run
        self.batch, self.pending, self.total = None, 0, 0

    def _op(self):
        if self.dry_run:
            self.total += 1
            return None
        if self.batch is None:
            self.batch = self.db.batch()
        self.pending += 1
        self.total += 1
        return self.batch

    def set(self, ref, data):
        b = self._op()
        if b is not None:
            b.set(ref, data)
            self._maybe_commit()

    def delete(self, ref):
        b = self._op()
        if b is not None:
            b.delete(ref)
            self._maybe_commit()

    def _maybe_commit(self):
        if self.pending >= BATCH_WRITE_LIMIT:
            self.flush()

    def flush(self):
        if self.batch is not None and self.pending:
            self.batch.commit()
        self.batch, self.pending = None, 0


def where_eq(query, field, value):
    from google.cloud.firestore_v1.base_query import FieldFilter
    return query.where(filter=FieldFilter(field, '==', value))


def trainee_id_number(tid):
    m = re.match(rf'^VIS-{YEAR}-(\d+)$', str(tid or ''))
    return int(m.group(1)) if m else 0


def find_users(db):
    admin = None
    for doc in where_eq(db.collection('users'), 'role', 'admin').stream():
        if doc.to_dict().get('isActive', True) is not False:
            admin = doc
            break
    if admin is None:
        raise SystemExit('No active user with role "admin" found in users.')
    trainers = [d for d in where_eq(db.collection('users'), 'role', 'trainer').stream()
                if d.to_dict().get('isActive', True) is not False]
    trainers.sort(key=lambda d: d.id)
    return admin, trainers


# ---------- seed ----------

def seed(db, excel_path, dry_run):
    from firebase_admin import firestore

    rng = random.Random(YEAR)
    trainees, warnings = read_trainees(excel_path)
    for w in warnings:
        print('WARN', w)
    if len(trainees) != EXPECTED_TRAINEES:
        print(f'WARN Found {len(trainees)} trainees in Excel (expected {EXPECTED_TRAINEES}).')

    admin, trainers = find_users(db)
    stamp = {'createdBy': admin.id, 'createdByRole': 'admin'}
    print(f'Admin: {admin.id}; trainers: {len(trainers)}')

    w = Writer(db, dry_run)
    counts = {k: 0 for k in ['batches', 'students', 'attendance', 'companies', 'placements']}
    skipped = {k: 0 for k in counts}

    # --- batches ---
    existing_batches = {}
    for doc in db.collection('batches').stream():
        name = doc.to_dict().get('batchName')
        if name:
            existing_batches.setdefault(name, doc.id)

    per_batch_count = {b: 0 for b in BATCHES}
    for t in trainees:
        per_batch_count[t['batch']] += 1

    batch_info = {}  # name -> dict(id, course, start, end)
    for i, (name, (course, start)) in enumerate(BATCHES.items()):
        end = add_months(start, COURSE_MONTHS[course])
        if name in existing_batches:
            skipped['batches'] += 1
            batch_info[name] = {'id': existing_batches[name], 'course': course, 'start': start, 'end': end}
            continue
        ref = db.collection('batches').document()
        trainer = trainers[i % len(trainers)] if trainers else None
        w.set(ref, {
            'batchName': name,
            'course': course,
            'trainerId': trainer.id if trainer else None,
            'trainerName': (trainer.to_dict().get('name') if trainer else None),
            'timing': BATCH_TIMING,
            'startDate': ist_midnight(start),
            'endDate': ist_midnight(end),
            'status': 'completed' if end < TODAY else 'active',
            'studentCount': per_batch_count[name],
            **stamp,
            'createdAt': firestore.SERVER_TIMESTAMP,
            'demoSeed': True,
        })
        counts['batches'] += 1
        batch_info[name] = {'id': ref.id, 'course': course, 'start': start, 'end': end}

    # --- students ---
    existing_students = {}
    max_existing_num = 0
    for doc in db.collection('students').stream():
        d = doc.to_dict()
        key = (str(d.get('name', '')).strip().lower(), norm_phone(d.get('phone')))
        existing_students.setdefault(key, (doc.id, d.get('traineeId')))
        max_existing_num = max(max_existing_num, trainee_id_number(d.get('traineeId')))

    counter_ref = db.collection('counters').document(COUNTER_DOC)
    counter_snap = counter_ref.get()
    counter_data = counter_snap.to_dict() if counter_snap.exists else {}
    current_count = int(counter_data.get('count', 0) or 0)
    base_count = counter_data.get('demoSeedBaseCount', current_count)
    next_num = max(current_count, max_existing_num) + 1

    seeded = []  # dicts with docId, traineeId + trainee fields
    for t in trainees:
        key = (t['name'].lower(), t['phone'])
        if key in existing_students:
            skipped['students'] += 1
            doc_id, tid = existing_students[key]
            seeded.append({**t, 'docId': doc_id, 'traineeId': tid, 'isNew': False})
            continue
        tid = f'VIS-{YEAR}-{next_num:04d}'
        next_num += 1
        ref = db.collection('students').document()
        b = batch_info[t['batch']]
        w.set(ref, {
            'traineeId': tid,
            'name': t['name'],
            'phone': t['phone'],
            'email': '',
            'gender': t['gender'],
            'dob': ist_midnight(t['dob']) if t['dob'] else None,
            'village': t['village'],
            'tehsil': t['tehsil'],
            'guardianName': '',
            'guardianPhone': t['guardianPhone'],
            'course': b['course'],
            'batch': t['batch'],
            'admissionDate': ist_midnight(b['start']),
            'status': t['status'],
            'totalFee': t['totalFee'],
            'feePaid': t['feePaid'],
            'feePending': t['totalFee'] - t['feePaid'],
            **stamp,
            'createdAt': firestore.SERVER_TIMESTAMP,
            'demoSeed': True,
        })
        counts['students'] += 1
        seeded.append({**t, 'docId': ref.id, 'traineeId': tid, 'isNew': True})
        existing_students[key] = (ref.id, tid)  # also skip duplicate rows within the Excel file

    final_count = next_num - 1
    if counts['students']:
        w.set(counter_ref, {
            'count': final_count,
            'updatedAt': firestore.SERVER_TIMESTAMP,
            'demoSeedBaseCount': base_count,
        })

    # --- attendance: each batch's own last 6 weeks (Mon-Sat), up to yesterday ---
    profiles = {}
    for s in seeded:
        if rng.random() < 0.08:
            profiles[s['docId']] = rng.uniform(0.45, 0.60)   # chronic absentee
        else:
            profiles[s['docId']] = rng.uniform(0.82, 0.96)

    present = marks = 0
    for name, b in batch_info.items():
        members = [s for s in seeded if s['batch'] == name and s['status'] not in NO_ATTENDANCE_STATUSES]
        if not members:
            continue
        last = min(b['end'], TODAY - timedelta(days=1))
        first = max(b['start'], last - timedelta(days=41))
        d = first
        while d <= last:
            if d.weekday() < 6:  # Mon-Sat
                doc_id = f"{b['id']}_{d.isoformat()}"
                ref = db.collection('attendance').document(doc_id)
                if ref.get().exists:
                    skipped['attendance'] += 1
                else:
                    status_map = {}
                    for s in members:
                        r = rng.random()
                        p = profiles[s['docId']]
                        if r < p:
                            status_map[s['docId']] = 'present'
                            present += 1
                        elif r < p + (1 - p) * 0.3:
                            status_map[s['docId']] = 'leave'
                        else:
                            status_map[s['docId']] = 'absent'
                        marks += 1
                    w.set(ref, {
                        'batchId': b['id'],
                        'batchName': name,
                        'date': ist_midnight(d),
                        'statusMap': status_map,
                        'markedAt': firestore.SERVER_TIMESTAMP,
                        'demoSeed': True,
                    })
                    counts['attendance'] += 1
            d += timedelta(days=1)

    # --- companies ---
    existing_companies = {str(doc.to_dict().get('companyName', '')).strip().lower()
                          for doc in db.collection('companies').stream()}
    for name, person, phone, loc, _ in COMPANIES:
        if name.lower() in existing_companies:
            skipped['companies'] += 1
            continue
        w.set(db.collection('companies').document(), {
            'companyName': name,
            'contactPerson': person,
            'contactNumber': phone,
            'location': loc,
            **stamp,
            'createdAt': firestore.SERVER_TIMESTAMP,
            'demoSeed': True,
        })
        counts['companies'] += 1

    # --- placements (Completed trainees only) ---
    placed_ids = {doc.to_dict().get('traineeId') for doc in db.collection('placements').stream()}
    completed = [s for s in seeded if s['status'] == 'Completed' and s['traineeId']]
    already = [s for s in completed if s['traineeId'] in placed_ids]
    skipped['placements'] = len(already)
    candidates = [s for s in completed if s['traineeId'] not in placed_ids]
    rng.shuffle(candidates)
    to_place = candidates[:max(0, TARGET_PLACEMENTS - len(already))]
    if len(completed) < TARGET_PLACEMENTS:
        print(f'WARN Only {len(completed)} Completed trainees, so only that many placements will be created.')

    for s in to_place:
        b = batch_info[s['batch']]
        course = b['course']
        options = [c for c in COMPANIES if course in c[4]]
        cname, person, _, loc, _ = rng.choice(options)
        joining = b['end'] + timedelta(days=rng.randint(7, 42))
        status = 'Joined' if rng.random() < 0.65 else 'Selected'
        if joining > TODAY:
            status = 'Selected'
        w.set(db.collection('placements').document(), {
            'studentName': s['name'],
            'traineeId': s['traineeId'],
            'course': course,
            'contactPersonName': person,
            'jobLocation': loc,
            'jobPost': rng.choice(JOB_POSTS[course]),
            'salary': str(rng.randrange(12000, 28001, 500)),
            'foodProvided': rng.choice(['Yes', 'No']),
            'companyName': cname,
            'joiningDate': ist_midnight(joining),
            'employmentType': 'Self-employed' if rng.random() < 0.2 else 'Salaried',
            'status': status,
            **stamp,
            'createdAt': firestore.SERVER_TIMESTAMP,
            'demoSeed': True,
        })
        counts['placements'] += 1

    w.flush()

    print()
    print('DRY RUN - nothing was written' if dry_run else 'Seed complete')
    for k in counts:
        print(f'  {k:<11} new: {counts[k]:>4}   skipped (existing): {skipped[k]:>4}')
    if counts['students']:
        print(f'  traineeId   VIS-{YEAR}-{final_count - counts["students"] + 1:04d} .. '
              f'VIS-{YEAR}-{final_count:04d}; counters/{COUNTER_DOC}.count = {final_count}')
    if marks:
        print(f'  attendance  present rate: {present / marks:.1%}')
    print(f'  total writes: {w.total}')


# ---------- clear ----------

def clear(db, dry_run):
    w = Writer(db, dry_run)
    for coll in SEED_COLLECTIONS:
        n = 0
        for doc in where_eq(db.collection(coll), 'demoSeed', True).stream():
            w.delete(doc.reference)
            n += 1
        w.flush()
        print(f'  {coll:<11} delete: {n}')

    counter_ref = db.collection('counters').document(COUNTER_DOC)
    snap = counter_ref.get()
    data = snap.to_dict() if snap.exists else {}
    if 'demoSeedBaseCount' in data:
        # So IDs of trainees created in the app after seeding don't collide
        remaining_max = 0
        for doc in db.collection('students').stream():
            d = doc.to_dict()
            if dry_run and d.get('demoSeed') is True:
                continue
            remaining_max = max(remaining_max, trainee_id_number(d.get('traineeId')))
        new_count = max(int(data['demoSeedBaseCount'] or 0), remaining_max)
        if not dry_run:
            from firebase_admin import firestore
            counter_ref.update({
                'count': new_count,
                'updatedAt': firestore.SERVER_TIMESTAMP,
                'demoSeedBaseCount': firestore.DELETE_FIELD,
            })
        print(f'  counters/{COUNTER_DOC}.count -> {new_count}')
    print('DRY RUN - nothing was deleted' if dry_run else 'Clear complete')


def main():
    ap = argparse.ArgumentParser(description='VTI ERP demo data seed')
    ap.add_argument('--excel', default=os.path.join(os.path.dirname(__file__), 'demo_trainees_Vistas.xlsx'),
                    help='path to demo_trainees_Vistas.xlsx')
    ap.add_argument('--dry-run', action='store_true', help='only print counts, write nothing')
    ap.add_argument('--clear', action='store_true', help='delete only records with demoSeed: true')
    args = ap.parse_args()

    if not args.clear and not os.path.isfile(args.excel):
        raise SystemExit(f'Excel file not found: {args.excel}')

    db = init_db()
    if args.clear:
        clear(db, args.dry_run)
    else:
        seed(db, args.excel, args.dry_run)


if __name__ == '__main__':
    sys.exit(main())
